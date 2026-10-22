"""Escritores de los formatos de Excel de destino.

Cada función recibe `filas` (lista de dicts del parser, con `monto_captura` y
`estado_reconc` ya agregados por el orquestador) y escribe un .xlsx.

    construir_aguas(filas, ruta)      -> hoja "Hoja1"      (formato único Aguas)
    construir_rendicion(filas, ruta)  -> hoja "para pegar"  (SC-020)
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from edificios import (CONFIG, CODIGO_CONTABLE, SERVICIO, es_numerico,
                       partir_cuenta, depto_consumo)

_ENC = Font(bold=True, color="FFFFFF")
_ENC_FILL = PatternFill("solid", fgColor="08424D")
_ROJO = PatternFill("solid", fgColor="F8D7DA")
_FMT_FECHA = "DD-MM-YYYY"
_FMT_MONTO = "#,##0"


def _encabezado(ws, cols, fila=1):
    for j, c in enumerate(cols, 1):
        cel = ws.cell(row=fila, column=j, value=c)
        cel.font = _ENC
        cel.fill = _ENC_FILL
        cel.alignment = Alignment(horizontal="center", vertical="center")


def _ajustar(ws, anchos):
    for j, a in enumerate(anchos, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = a


def _marcar_dif(celdas, fila):
    """Pinta de rojo las celdas de monto cuando la reconciliación no coincide."""
    if fila.get("estado_reconc") == "DIFERENCIA":
        for c in celdas:
            c.fill = _ROJO


def _num(f):
    try:
        return int(str(f.get("depto", "0")) or 0)
    except ValueError:
        return 0


def _orden_num_primero(filas):
    """Deptos numéricos ascendentes y las unidades especiales al final."""
    return sorted(filas, key=lambda f: (0, _num(f)) if es_numerico(f.get("depto"))
                  else (1, 0, str(f.get("depto")).upper()))


# ── Aguas (formato único tipo Viento Sur) ────────────────────────────────────────
def construir_aguas(filas, ruta):
    cfg = CONFIG["aguas"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = cfg["hoja"]
    cols = ["EMPRESA", "N° CLIENTE", "N° CLIENTE2", "DIGITO VERIFICADOR", "DEPTO",
            "FECHA CHECK IN", "TORRE", "Monto en boletas", "Monto en capturas"]
    _encabezado(ws, cols)
    r = 2
    total_b = 0
    for f in _orden_num_primero(filas):
        n_cta, cuerpo, dv = partir_cuenta(f.get("n_cuenta", ""))
        ws.cell(row=r, column=1, value=f.get("empresa", ""))
        ws.cell(row=r, column=2, value=n_cta)
        ws.cell(row=r, column=3, value=cuerpo)
        ws.cell(row=r, column=4, value=dv)
        ws.cell(row=r, column=5, value=depto_consumo(f["depto"]))
        cf = ws.cell(row=r, column=6, value=f.get("vencimiento"))
        cf.number_format = _FMT_FECHA
        ws.cell(row=r, column=7, value=f.get("torre", ""))
        cb = ws.cell(row=r, column=8, value=f.get("monto"))
        cp = ws.cell(row=r, column=9, value=f.get("monto_captura"))
        cb.number_format = cp.number_format = _FMT_MONTO
        _marcar_dif([cb, cp], f)
        if isinstance(f.get("monto"), int):
            total_b += f["monto"]
        r += 1
    ws.cell(row=r, column=7, value="Total:").font = Font(bold=True)
    ct = ws.cell(row=r, column=8, value=total_b)
    ct.number_format = _FMT_MONTO
    ct.font = Font(bold=True)
    _ajustar(ws, [14, 13, 13, 10, 14, 15, 7, 16, 16])
    wb.save(ruta)


# ── Rendición SC-020 (Enel / Aguas) — hoja "para pegar" ──────────────────────────
def _detalle(f):
    serv = SERVICIO.get(f.get("proveedor"), "")
    depto = f.get("depto", "")
    pini, pfin = f.get("periodo", (None, None))
    if pini and pfin:
        per = f" - Periodo {pini.strftime('%d-%m')} al {pfin.strftime('%d-%m')}"
    else:
        per = ""
    return f"{serv} - Depto {depto}{per}".strip()


def construir_rendicion(filas, ruta):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = CONFIG["rendicion"]["hoja"]
    cols = ["FECHA", "PROVEEDOR", "TIPO DOCUMENTO", "N° DOCUMENTO",
            "DETALLE", "MONTO", "CODIGO CONTABLE"]
    _encabezado(ws, cols)
    r = 2
    for f in _orden_num_primero(filas):
        cf = ws.cell(row=r, column=1, value=f.get("fecha_emision"))
        cf.number_format = _FMT_FECHA
        ws.cell(row=r, column=2, value=(f.get("proveedor") or "").title())
        ws.cell(row=r, column=3, value=f.get("tipo_doc", ""))
        doc = f.get("folio", "")
        ws.cell(row=r, column=4, value=int(doc) if str(doc).isdigit() else doc)
        ws.cell(row=r, column=5, value=_detalle(f))
        cm = ws.cell(row=r, column=6, value=f.get("monto"))
        cm.number_format = _FMT_MONTO
        ws.cell(row=r, column=7, value=CODIGO_CONTABLE.get(f.get("proveedor"), ""))
        r += 1
    _ajustar(ws, [12, 12, 15, 14, 42, 12, 24])
    wb.save(ruta)


ESCRITORES = {
    "aguas": construir_aguas,
    "rendicion": construir_rendicion,
}
